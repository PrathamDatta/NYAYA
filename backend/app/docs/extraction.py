"""Text extraction for PDFs, DOCX, plain text and (optionally) images. Pure-Python, no external binaries."""
from __future__ import annotations

import io
import re
from dataclasses import dataclass
from pathlib import Path


class ExtractionError(RuntimeError):
    pass


class OcrUnavailableError(ExtractionError):
    pass


@dataclass
class Page:
    page: int
    text: str
    method: str  # text | ocr


_SPACES = re.compile(r"[ \t ]+")


def _norm(t: str) -> str:
    t = t.replace("\x00", "")
    t = re.sub(r"(\w)-\n(\w)", r"\1\2", t)  # de-hyphenate line-end breaks
    return _SPACES.sub(" ", t).strip()


def extract_pdf(data: bytes, max_pages: int = 300) -> list[Page]:
    from pdfminer.high_level import extract_text
    from pdfminer.pdfparser import PDFSyntaxError

    try:
        from pypdf import PdfReader

        n = len(PdfReader(io.BytesIO(data)).pages)
    except Exception as e:  # noqa: BLE001
        raise ExtractionError(f"Unreadable PDF: {type(e).__name__}") from e
    if n > max_pages:
        raise ExtractionError(f"PDF has {n} pages; the limit is {max_pages}")
    try:
        raw = extract_text(io.BytesIO(data))
    except (PDFSyntaxError, ValueError, TypeError) as e:
        raise ExtractionError(f"Unreadable PDF: {type(e).__name__}") from e
    parts = raw.split("\f")
    if parts and not parts[-1].strip():
        parts = parts[:-1]
    pages = [Page(i + 1, _norm(p), "text") for i, p in enumerate(parts)]
    # pad if the extractor produced fewer pages than the document has (blank pages)
    while len(pages) < n:
        pages.append(Page(len(pages) + 1, "", "text"))
    return pages


def extract_docx(data: bytes) -> list[Page]:
    import docx

    try:
        d = docx.Document(io.BytesIO(data))
    except Exception as e:  # noqa: BLE001
        raise ExtractionError(f"Unreadable DOCX: {type(e).__name__}") from e
    paras = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            paras.append(" | ".join(c.text.strip() for c in row.cells))
    return [Page(1, _norm("\n".join(paras)), "text")]


def extract_text_file(data: bytes) -> list[Page]:
    for enc in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            return [Page(1, _norm(data.decode(enc)), "text")]
        except UnicodeError:
            continue
    raise ExtractionError("Could not decode text file")


def extract_image(data: bytes) -> list[Page]:
    try:
        from rapidocr_onnxruntime import RapidOCR  # optional dependency: pip install rapidocr-onnxruntime
    except ImportError as e:
        raise OcrUnavailableError(
            "OCR is not installed. Install the optional package with `pip install rapidocr-onnxruntime` "
            "to read scanned images.") from e
    import numpy as np
    from PIL import Image

    try:
        img = Image.open(io.BytesIO(data)).convert("RGB")
    except Exception as e:  # noqa: BLE001
        raise ExtractionError("Unreadable image") from e
    result, _ = RapidOCR()(np.array(img))
    text = "\n".join(r[1] for r in (result or []))
    return [Page(1, _norm(text), "ocr")]


def sniff(data: bytes, filename: str) -> str:
    """Return a canonical kind from magic bytes (never from the client-supplied content type)."""
    ext = Path(filename).suffix.lower()
    if data[:5] == b"%PDF-":
        return "pdf"
    if data[:4] == b"PK\x03\x04" and ext == ".docx":
        import zipfile

        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                names = z.namelist()
                if "word/document.xml" not in names:
                    raise ExtractionError("Not a Word document")
                if sum(i.file_size for i in z.infolist()) > 200 * 1024 * 1024:
                    raise ExtractionError("Document expands too large")
        except zipfile.BadZipFile as e:
            raise ExtractionError("Corrupt DOCX") from e
        return "docx"
    if data[:8] == b"\x89PNG\r\n\x1a\n" or data[:3] == b"\xff\xd8\xff":
        return "image"
    if ext in (".txt", ".md") and b"\x00" not in data[:4096]:
        return "text"
    raise ExtractionError("Unsupported file type. Supported: PDF, DOCX, TXT, PNG/JPEG (OCR if installed).")


def extract(data: bytes, kind: str, max_pages: int = 300) -> list[Page]:
    if kind == "pdf":
        return extract_pdf(data, max_pages)
    if kind == "docx":
        return extract_docx(data)
    if kind == "text":
        return extract_text_file(data)
    if kind == "image":
        return extract_image(data)
    raise ExtractionError("Unsupported file type")
